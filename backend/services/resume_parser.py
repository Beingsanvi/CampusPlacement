"""
Resume parsing - extract plain text from PDF, DOCX, or TXT resumes.
"""

import logging
import os

logger = logging.getLogger(__name__)


def extract_text(file_path: str) -> str:
    """
    Extract raw text from a resume file (.pdf, .docx, .txt).

    Args:
        file_path: Absolute path to the resume file

    Returns:
        Extracted plain text (utf-8)

    Raises:
        ValueError: if the file type is unsupported or not found
    """
    if not os.path.exists(file_path):
        raise ValueError(f"File not found: {file_path}")

    ext = os.path.splitext(file_path)[1].lower()

    try:
        if ext == ".txt":
            return _read_txt(file_path)
        if ext == ".pdf":
            return _read_pdf(file_path)
        if ext == ".docx":
            return _read_docx(file_path)
        if ext in (".doc", ".odt"):
            raise ValueError(
                f"{ext.upper()} files are not supported. Please upload a PDF, DOCX, or TXT file."
            )
        raise ValueError(f"Unsupported file type: {ext}")
    except ValueError:
        raise
    except Exception as e:
        logger.error("Failed to parse %s: %s", file_path, e)
        raise ValueError(f"Could not parse the file: {e}")


def _read_txt(path: str) -> str:
    for encoding in ("utf-8", "latin-1"):
        try:
            with open(path, "r", encoding=encoding) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    with open(path, "r", encoding="latin-1") as f:
        return f.read()


def _read_pdf(path: str) -> str:
    from pypdf import PdfReader

    reader = PdfReader(path)
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(pages)


def _read_docx(path: str) -> str:
    import docx

    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text.strip())
    return "\n".join(parts)